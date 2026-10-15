"""
Analysis-Ready Export Builder
=============================

Purpose:
    Turn the boq_coords knowledge bank into one analysis-ready Excel
    workbook: non-analysis columns moved out of the main sheet,
    currency converted to INR beside the original amount, dates,
    units, statuses and placeholders standardized, and the BoQ item
    hierarchy made explicit so document totals are never double-counted.

    Product name, make and model are carried through UNCHANGED
    (product_name, description, make_canonical, model_canonical) - their
    standardization waits for the user's own mapping tables.

Input:
    Quotation_Data/07_knowledge_bank_coords/knowledge_bank_items_coords_<date>.csv
    (newest dated file by default; override with --input). Read-only.

Output (new dated files, never overwritten - a _v2/_v3 suffix is added
if today's name already exists):
    Quotation_Data/08_analysis_ready/knowledge_bank_analysis_<date>.xlsx
        Items           one row per BoQ line (the analysis sheet)
        Quotations      one row per document, order-level money lives here
        Audit           doc_id + line_seq + the source columns moved off Items
        FX_Rates        the rate table used for *_inr columns
        Data_Dictionary every column on Items / Quotations explained
    Quotation_Data/08_analysis_ready/analysis_export_validation_<date>.txt

Design:
    - Raw values stay beside derived ones: amounts keep their original
      currency (*_orig) next to *_inr, with fx_rate_to_inr / fx_source
      recording how each INR figure was obtained.
    - One blank convention: an empty cell means unknown / not applicable.
      Placeholder text (Unknown, NA, -) and "none" zeros (PO value 0)
      become empty; an accompanying status column says why.
    - Order-level Odoo money (PO value, price in INR, potential) appears
      only on Quotations, where it can be summed over
      is_latest_revision == Y without inflating by line count.
    - Nothing is dropped: every source row appears on Items and Audit;
      exact duplicate rows are flagged (duplicate_line), not removed.

Important:
    The workbook contains customer names - it is confidential and lives
    under the gitignored Quotation_Data/. Console output is aggregate
    counts only; still, redirect stdout to a log file when running.

This script does NOT modify its input file.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parent.parent


# ============================================================
# CONFIGURATION
# ============================================================

INPUT_FOLDER = PROJECT_ROOT / "Quotation_Data" / "07_knowledge_bank_coords"
INPUT_NAME_PATTERN = re.compile(r"^knowledge_bank_items_coords_\d{4}-\d{2}-\d{2}\.csv$")

OUTPUT_FOLDER = PROJECT_ROOT / "Quotation_Data" / "08_analysis_ready"

BASE_CURRENCY = "INR"

# Odoo's own fixed conversion rates - identical on every order in the
# export (measured 2026-09-28: matched_order_price_in_inr divided by
# matched_order_latest_price_quoted), so INR totals here reconcile with
# Odoo's. They ignore exchange-rate movement over 2024-26.
FX_RATES_TO_INR = {"USD": 84.84, "EUR": 94.25, "SAR": 22.61}
FX_SOURCE_BASE = "BASE_CURRENCY"
FX_SOURCE_ODOO = "ODOO_FIXED_RATE"
FX_SOURCE_NO_RATE = "NO_RATE"

PLACEHOLDERS = {"UNKNOWN", "NA", "N/A", "-", "--", "NIL", "NULL"}

UNIT_STD = {
    "NOS": "NOS", "EACH": "NOS", "UNIT": "NOS", "PCS": "NOS",
    "SET": "SET", "LOT": "LOT", "JOB": "JOB",
    "METER": "METER",
    "DAY": "DAY",
    "MAN-DAY": "MAN-DAY", "MAN-DAYS": "MAN-DAY", "MANDAY": "MAN-DAY",
    "MANDAYS": "MAN-DAY", "MDY": "MAN-DAY",
    "VISIT": "VISIT", "VISITS": "VISIT",
    "MONTH": "MONTH", "YEAR": "YEAR",
}

# A unit price is only comparable with another in the same class. VISIT
# is its own class: a per-visit service charge is neither a count of
# goods nor a duration.
UNIT_CLASS = {
    "NOS": "COUNT",
    "SET": "BUNDLE", "LOT": "BUNDLE", "JOB": "BUNDLE",
    "METER": "LENGTH",
    "DAY": "TIME", "MAN-DAY": "TIME", "MONTH": "TIME", "YEAR": "TIME",
    "VISIT": "SERVICE_VISIT",
}

QUOTE_STATUS = {"QUOTE - YES": "YES", "QUOTE - PENDING": "PENDING", "QUOTE - REGRET": "REGRET"}

OUTCOME_GROUP = {
    "WON": "WON",
    "LOST": "LOST",
    "DROP": "DROPPED",
    "ACTIVE": "OPEN",
    "HOLD": "OPEN",
    "SYSTEM CLOSED": "CLOSED_NO_OUTCOME",
}

MATCH_QUALITY_REVIEW = {"REVIEW", "LOW", "RFQ_AMBIGUOUS"}

STATE_FIXES = {
    "Chattisgarh": "Chhattisgarh",
    "กรุงเทพมหานคร": "Bangkok",
}

MONTHS = {
    name: number
    for number, names in enumerate(
        [("JAN", "JANUARY"), ("FEB", "FEBRUARY"), ("MAR", "MARCH"),
         ("APR", "APRIL"), ("MAY",), ("JUN", "JUNE"), ("JUL", "JULY"),
         ("AUG", "AUGUST"), ("SEP", "SEPT", "SEPTEMBER"), ("OCT", "OCTOBER"),
         ("NOV", "NOVEMBER"), ("DEC", "DECEMBER")],
        start=1,
    )
    for name in names
}

TOTAL_ROW_RX = re.compile(r"^\s*(?:grand\s+|sub[\s-]*)?total\b", re.IGNORECASE)
REVISION_RX = re.compile(r"^(?P<base>.*\d.*?)[\s_\-]*R(?P<rev>\d{1,2})$", re.IGNORECASE)
NUMERIC_DATE_RX = re.compile(r"^\s*(\d{1,2})\s*[./-]\s*(\d{1,2})\s*[./-]\s*(\d{2,4})\s*$")
TEXT_DATE_RX = re.compile(r"^\s*(\d{1,2})(?:st|nd|rd|th)?[\s,.-]+([A-Za-z]+)[\s,.-]+(\d{4})\s*$")
EXCEL_ILLEGAL_RX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
EXCEL_CELL_LIMIT = 32767


# ============================================================
# VALUE HELPERS (pure - unit-tested in tests/test_analysis_export.py)
# ============================================================

def clean_text(value) -> str:
    """Trim whitespace; placeholder text becomes empty."""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return "" if text.upper() in PLACEHOLDERS else text


def is_text(series: pd.Series) -> bool:
    return series.dtype == object or pd.api.types.is_string_dtype(series)


def to_number(value):
    text = str(value or "").strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def doc_id_for(source_path: str) -> str:
    """Short, stable document key - the same path always gives the same id."""
    return "D" + hashlib.sha1(source_path.encode("utf-8")).hexdigest()[:8].upper()


def parse_revision(quotation_number: str) -> tuple[str, int]:
    """'Q25N10072R2' -> ('Q25N10072', 2). No R-suffix -> (number, 0)."""
    text = clean_text(quotation_number)
    match = REVISION_RX.match(text)
    if match:
        return match.group("base").rstrip(" _-"), int(match.group("rev"))
    return text, 0


def parse_document_date(raw: str):
    """Day-first parse of a date printed on a document; None if unreadable."""
    text = clean_text(raw)
    match = NUMERIC_DATE_RX.match(text)
    if match:
        day, month, year = (int(g) for g in match.groups())
        if year < 100:
            year += 2000
    else:
        match = TEXT_DATE_RX.match(text)
        if not match or match.group(2).upper() not in MONTHS:
            return None
        day, month, year = int(match.group(1)), MONTHS[match.group(2).upper()], int(match.group(3))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def fiscal_year(value) -> str:
    """Indian fiscal year (April-March): 2025-06-23 -> 'FY2025-26'."""
    if value is None or pd.isna(value):
        return ""
    start = value.year if value.month >= 4 else value.year - 1
    return f"FY{start}-{str(start + 1)[-2:]}"


def standardize_unit(unit_normalized: str) -> str:
    key = clean_text(unit_normalized).upper()
    return UNIT_STD.get(key, key)


def standardize_state(state: str) -> str:
    text = re.sub(r"\s*\([A-Z]{2}\)\s*$", "", clean_text(state))
    return STATE_FIXES.get(text, text)


def standardize_regions(regions: str) -> str:
    parts = {clean_text(part) for part in str(regions or "").split(",")}
    return "|".join(sorted(part for part in parts if part))


def parse_percent(value):
    match = re.match(r"^\s*(\d+(?:\.\d+)?)\s*%\s*$", str(value or ""))
    return float(match.group(1)) if match else None


def parse_days(value):
    match = re.match(r"^\s*(\d+(?:\.\d+)?)\s*days?\s*$", str(value or ""), re.IGNORECASE)
    return float(match.group(1)) if match else None


def expected_finalization(month: str, year: str) -> str:
    """('September', '2026') -> '2026-09'; blank if either is Unknown."""
    month_no = MONTHS.get(clean_text(month).upper())
    year_text = clean_text(year)
    if not month_no or not re.fullmatch(r"\d{4}", year_text):
        return ""
    return f"{year_text}-{month_no:02d}"


def match_quality(status: str, customer_id: str) -> str:
    if not clean_text(customer_id):
        return "UNMATCHED"
    return "REVIEW" if clean_text(status).upper() in MATCH_QUALITY_REVIEW else "MATCHED"


def normalize_item_no(item_no: str) -> str:
    return str(item_no or "").strip().rstrip(".").strip().upper()


def classify_lines(levels, quantities, price_statuses, names, item_nos, parent_item_nos):
    """
    One document's rows, in document order -> (line_type, parent_line_seq,
    bundle_seq) per row, with line_seq = position + 1.

    Level-0 rows (no item number) are HEADER only when they carry neither
    a quantity nor any price information; otherwise they are real,
    unnumbered lines.

    A sub-item's parent is the nearest earlier row whose item_no equals
    its parsed parent_item_no. Position alone is not used: variants such
    as 1A, 1B are often printed with no row "1" at all, and linking them
    to whatever numbered row came before would attach 5A/5B to item 4.
    When the parent is not printed, parent_line_seq stays None and
    consecutive siblings sharing a parent_item_no share one bundle.
    """
    types, parents, bundles = [], [], []
    for index, level in enumerate(levels):
        seq = index + 1
        if TOTAL_ROW_RX.match(names[index] or "") and not item_nos[index]:
            line_type = "TOTAL"
        elif level == 0:
            is_header = not quantities[index] and price_statuses[index] == "MISSING"
            line_type = "HEADER" if is_header else "LINE"
        elif level == 1:
            line_type = "LINE"
        else:
            line_type = "SUB_LINE"

        parent = None
        wanted = normalize_item_no(parent_item_nos[index])
        if line_type == "SUB_LINE" and wanted:
            for back in range(index - 1, -1, -1):
                if normalize_item_no(item_nos[back]) == wanted and levels[back] < level:
                    parent = back + 1
                    break

        if line_type in ("HEADER", "TOTAL"):
            bundle = None
        elif parent is not None and bundles[parent - 1] is not None:
            bundle = bundles[parent - 1]
        elif (line_type == "SUB_LINE" and parent is None and index > 0
              and types[index - 1] == "SUB_LINE" and parents[index - 1] is None
              and normalize_item_no(parent_item_nos[index - 1]) == wanted):
            bundle = bundles[index - 1]
        else:
            bundle = seq

        types.append(line_type)
        parents.append(parent)
        bundles.append(bundle)
    return types, parents, bundles


def include_in_total(line_types, parents, usable):
    """
    'Y' when a row's total should count toward the document total: its
    price is usable and no ancestor's price is already counted (a priced
    parent already includes its children).
    """
    counted = []
    for index, line_type in enumerate(line_types):
        ok = usable[index] and line_type != "TOTAL"
        parent = parents[index]
        while ok and parent is not None:
            if counted[parent - 1]:
                ok = False
            parent = parents[parent - 1]
        counted.append(ok)
    return ["Y" if flag else "N" for flag in counted]


# ============================================================
# COLUMN LAYOUT
# ============================================================

ITEM_COLUMNS = [
    # quotation
    "doc_id", "line_seq", "quotation_number", "quotation_base_no", "revision_no",
    "revision_group", "is_latest_revision", "odoo_rfq_number", "odoo_order_id",
    "quotation_date", "document_date", "quotation_year", "quotation_month",
    "fiscal_year", "date_source", "date_ambiguous",
    # customer
    "customer_id", "customer_name", "customer_type", "end_user_id", "end_user_name",
    "industry", "industry_detail", "industry_source", "city", "state", "country",
    "sales_regions", "customer_match_quality",
    # item
    "item_no", "parent_item_no", "item_level", "line_type", "parent_line_seq",
    "bundle_id", "product_name", "description", "product_family",
    "make_canonical", "model_canonical", "quantity", "unit_std", "unit_class",
    # pricing
    "price_status", "currency_orig", "currency_source", "unit_price_orig",
    "total_price_orig", "fx_rate_to_inr", "fx_source", "unit_price_inr",
    "total_price_inr", "price_basis", "price_quality", "item_confidence",
    "price_usable", "include_in_doc_total", "duplicate_line",
    # outcome (no money - that's on Quotations)
    "quote_status", "quote_nature", "outcome_status", "outcome_group",
    "winning_chance_pct", "sbu", "spares_type", "service_type",
]

QUOTATION_COLUMNS = [
    "doc_id", "quotation_number", "quotation_base_no", "revision_no",
    "revision_group", "is_latest_revision", "odoo_rfq_number", "odoo_order_id",
    "quotation_date", "document_date", "quotation_year", "quotation_month",
    "fiscal_year", "date_source", "date_ambiguous",
    "customer_id", "customer_name", "customer_type", "end_user_id", "end_user_name",
    "industry", "industry_detail", "industry_source", "city", "state", "country",
    "sales_regions", "customer_match_quality",
    "quote_status", "quote_nature", "type_of_quote", "outcome_status",
    "outcome_group", "winning_chance_pct", "sbu", "spares_type", "service_type",
    "expected_finalization", "cycle_time_days", "loss_reason", "loss_reason_detail",
    "po_number", "order_currency", "order_latest_price_quoted", "order_price_inr",
    "order_po_value_orig", "order_po_value_inr", "order_potential_estimate_inr",
    "line_count", "priced_line_count", "doc_currency", "sum_items_orig",
    "sum_items_inr", "reconcile_ratio_vs_odoo",
]

# Source columns copied verbatim to Audit (everything not carried onto
# Items unchanged). source_file/source_path can contain customer names.
AUDIT_SOURCE_COLUMNS = [
    "source_file", "source_path", "data_source", "quotation_date_raw",
    "order_date_raw", "matched_order_create_date", "order_state",
    "matched_order_rfq_reference_number", "matched_order_control_no",
    "matched_industry", "matched_industry_specify_others",
    "matched_order_industry", "matched_order_industry_specify_others",
    "matched_industry_confidence", "customer_match_status", "customer_match_score",
    "make", "model", "make_normalized", "model_normalized",
    "make_canonical_basis", "model_canonical_basis", "product_family_basis",
    "unit", "unit_normalized", "unit_class",
    "unit_price", "total_price", "unit_price_raw", "total_price_raw",
    "total_price_source", "unit_price_final", "total_price_final",
    "currency", "document_currency", "arithmetic_check", "validation_error",
]

ORDER_MONEY_COLUMNS = {
    "order_latest_price_quoted", "order_price_inr", "order_po_value_orig",
    "order_po_value_inr", "order_potential_estimate_inr",
}

DATE_COLUMNS = {"quotation_date", "document_date"}
MONEY_COLUMNS = {
    "unit_price_orig", "total_price_orig", "unit_price_inr", "total_price_inr",
    "order_price_inr", "order_po_value_orig", "order_po_value_inr",
    "order_potential_estimate_inr", "sum_items_orig", "sum_items_inr",
}

ALLOWED_VALUES = {
    "line_type": {"LINE", "SUB_LINE", "HEADER", "TOTAL"},
    "price_usable": {"Y", "N"},
    "include_in_doc_total": {"Y", "N"},
    "duplicate_line": {"Y", "N"},
    "is_latest_revision": {"Y", "N"},
    "customer_match_quality": {"MATCHED", "REVIEW", "UNMATCHED"},
    "industry_source": {"ODOO_ORDER", "CUSTOMER_HISTORY", ""},
    "fx_source": {FX_SOURCE_BASE, FX_SOURCE_ODOO, FX_SOURCE_NO_RATE, ""},
    "quote_status": set(QUOTE_STATUS.values()) | {""},
    "quote_nature": {"FIRM", "BUDGETARY", "TENDER", "ARC", ""},
    "outcome_group": set(OUTCOME_GROUP.values()) | {""},
    "unit_std": set(UNIT_STD.values()) | {""},
    "unit_class": set(UNIT_CLASS.values()) | {""},
    "price_status": {"NUMERIC", "QUOTED_SEPARATELY", "INCLUDED", "MISSING"},
    "currency_source": {"CELL", "DOCUMENT", "DEFAULT", ""},
}

DICTIONARY = {
    "doc_id": "Short stable id of the source document (hash of source_path). Join key between all sheets.",
    "line_seq": "Row position within the document, in document order (1 = first row).",
    "quotation_number": "Quotation number as parsed from the document. Not unique; junk values possible.",
    "quotation_base_no": "quotation_number without its revision suffix (Q25N10072R2 -> Q25N10072).",
    "revision_no": "Revision from the R-suffix of quotation_number; 0 when there is none.",
    "revision_group": "Documents that are revisions of one another: the Odoo order id, else Q:<quotation_base_no>, else doc_id.",
    "is_latest_revision": "Y on the document with the highest revision_no (then latest document_date) in its revision_group. Use for win/loss and current price levels; use all rows to see price movement between revisions.",
    "odoo_rfq_number": "Odoo RFQ assignment number of the matched order.",
    "odoo_order_id": "Odoo sale order id. Blank = no Odoo order resolved.",
    "quotation_date": "Headline date: Odoo order creation date when matched, else the date printed on the document. Revisions of one order share it.",
    "document_date": "Date printed on the document itself, parsed day-first. Use it to order revisions.",
    "quotation_year": "Calendar year of quotation_date.",
    "quotation_month": "YYYY-MM of quotation_date.",
    "fiscal_year": "Indian fiscal year (April-March) of quotation_date, e.g. FY2025-26.",
    "date_source": "Where quotation_date came from: ODOO_ORDER (reliable), PARSED_DOC (regex from PDF), NONE.",
    "date_ambiguous": "YES when day and month were both <= 12 and day-first was assumed - treat as approximate.",
    "customer_id": "Odoo customer id. Group customers by this, not by name (some ids share a name).",
    "customer_name": "Odoo customer name (Odoo spelling, trimmed). CONFIDENTIAL.",
    "customer_type": "END USER / EPC / EPMC.",
    "end_user_id": "Odoo end-user id (on EPC quotations the end user is the real customer).",
    "end_user_name": "Odoo end-user name. CONFIDENTIAL.",
    "industry": "Industry recorded on the Odoo order; falls back to the customer's usual industry.",
    "industry_detail": "Free-text detail, only when industry is Others. Uppercased.",
    "industry_source": "ODOO_ORDER (from this order) or CUSTOMER_HISTORY (customer-level estimate).",
    "city": "Customer city (Odoo).",
    "state": "Customer state/province, without the (IN)-style country suffix.",
    "country": "Customer country (Odoo).",
    "sales_regions": "Sales-team region tags, sorted and pipe-joined. Not geography.",
    "customer_match_quality": "MATCHED / REVIEW (weak or ambiguous match - verify) / UNMATCHED (no customer).",
    "item_no": "Item number as printed. Repeats within a document (separate schedules).",
    "parent_item_no": "Parent item number as parsed, for sub-items.",
    "item_level": "0 = no item number, 1 = main item, 2-3 = sub-item.",
    "line_type": "LINE (item), SUB_LINE (child of an item), HEADER (section heading, no qty/price), TOTAL (total row - never sum it).",
    "parent_line_seq": "line_seq of this row's parent row (the earlier row whose item_no = parent_item_no). Blank when the parent is not printed as its own row (1A, 1B with no row 1).",
    "bundle_id": "doc_id-line_seq of the top-level row this row belongs to. Groups an item with its sub-items; sub-items whose parent is not printed are grouped with their consecutive siblings.",
    "product_name": "Product name as extracted. Not yet standardized.",
    "description": "Full description as extracted. Not yet standardized.",
    "product_family": "Rule-based product family. Blank is expected for most rows (no rule matched), not junk.",
    "make_canonical": "Make with spelling variants merged (SIEMENS AG, GERMANY -> SIEMENS). Alternates pipe-joined.",
    "model_canonical": "Model with cosmetic variants merged. OXYMAT 61 and 64 stay different.",
    "quantity": "Quantity as printed. Blank = not printed (never assumed to be 1).",
    "unit_std": "Standardized unit: NOS, SET, LOT, JOB, METER, DAY, MAN-DAY, VISIT, MONTH, YEAR.",
    "unit_class": "COUNT / BUNDLE / LENGTH / TIME / SERVICE_VISIT. Compare unit prices only within one class.",
    "price_status": "NUMERIC (a number) / QUOTED_SEPARATELY / INCLUDED (in another line - not zero) / MISSING.",
    "currency_orig": "Currency of the printed amounts.",
    "currency_source": "How currency_orig was decided: CELL (printed in the price cell), DOCUMENT (currency markers on the table pages), DEFAULT (INR assumed - unconfirmed).",
    "unit_price_orig": "Unit price in currency_orig, as quoted.",
    "total_price_orig": "Line total in currency_orig, as quoted.",
    "fx_rate_to_inr": "Rate applied to get the *_inr columns (1 for INR).",
    "fx_source": "BASE_CURRENCY (already INR), ODOO_FIXED_RATE (Odoo's fixed rate, see FX_Rates), NO_RATE (INR left blank).",
    "unit_price_inr": "unit_price_orig x fx_rate_to_inr.",
    "total_price_inr": "total_price_orig x fx_rate_to_inr.",
    "price_basis": "REPORTED (printed), DERIVED_FROM_TOTAL (unit = total / qty), NONE (no price).",
    "price_quality": "TRUSTED / FLAGGED (failed a sanity check) / NON_NUMERIC.",
    "item_confidence": "HIGH / LOW extraction confidence. Segment by it rather than averaging blind.",
    "price_usable": "Y when the price is safe for price analysis: TRUSTED, REPORTED or DERIVED_FROM_TOTAL, above zero, not a TOTAL row.",
    "include_in_doc_total": "Y when the line total should be summed into a document total (price usable and no priced parent already covers it).",
    "duplicate_line": "Y on the 2nd+ copy of a row identical to an earlier one in the source. Reviewed as genuine repeats in the documents - kept.",
    "quote_status": "YES / PENDING / REGRET (Odoo).",
    "quote_nature": "FIRM / BUDGETARY / TENDER / ARC (Odoo).",
    "outcome_status": "Odoo present status: WON, LOST, DROP, ACTIVE, HOLD, SYSTEM CLOSED.",
    "outcome_group": "WON / LOST / DROPPED / OPEN (ACTIVE, HOLD) / CLOSED_NO_OUTCOME (SYSTEM CLOSED).",
    "winning_chance_pct": "Odoo winning chance as a number (25/50/75/100).",
    "sbu": "Odoo SBU type.",
    "spares_type": "Odoo spares type(s), pipe-joined.",
    "service_type": "Odoo service type.",
    "type_of_quote": "Techno - Commercial / Technical (Odoo).",
    "expected_finalization": "YYYY-MM from Odoo's tentative finalization month and year.",
    "cycle_time_days": "Odoo average cycle time, in days.",
    "loss_reason": "Odoo structured reason for loss.",
    "loss_reason_detail": "Odoo free-text main reason of losing the order.",
    "po_number": "Customer PO number (Odoo).",
    "order_currency": "Currency of the Odoo order.",
    "order_latest_price_quoted": "Odoo's latest price quoted, as text with its currency symbol. ORDER-LEVEL.",
    "order_price_inr": "Odoo's quoted price in INR. ORDER-LEVEL: sum only over is_latest_revision = Y.",
    "order_po_value_orig": "PO value in the order currency; blank = no PO. ORDER-LEVEL.",
    "order_po_value_inr": "PO value in INR; blank = no PO. ORDER-LEVEL.",
    "order_potential_estimate_inr": "Odoo total potential estimate; blank = not estimated. ORDER-LEVEL.",
    "line_count": "Rows on Items for this document.",
    "priced_line_count": "Rows with include_in_doc_total = Y.",
    "doc_currency": "The document's item currency; MIXED if lines use more than one.",
    "sum_items_orig": "Sum of included line totals in doc_currency (blank when MIXED).",
    "sum_items_inr": "Sum of included line totals in INR.",
    "reconcile_ratio_vs_odoo": "sum_items_inr / order_price_inr. Near 1 = the document's lines match Odoo's quoted total.",
}


# ============================================================
# BUILD
# ============================================================

def find_default_input() -> Path | None:
    candidates = sorted(
        path for path in INPUT_FOLDER.glob("knowledge_bank_items_coords_*.csv")
        if INPUT_NAME_PATTERN.match(path.name)
    )
    return candidates[-1] if candidates else None


def unique_output_path(folder: Path, stem: str, suffix: str) -> Path:
    path = folder / f"{stem}{suffix}"
    version = 2
    while path.exists():
        path = folder / f"{stem}_v{version}{suffix}"
        version += 1
    return path


def build_items(src: pd.DataFrame) -> pd.DataFrame:
    out = pd.DataFrame(index=src.index)
    t = src.map(clean_text)

    # --- quotation ---
    out["doc_id"] = src["source_path"].map(doc_id_for)
    out["line_seq"] = src.groupby("source_path").cumcount() + 1
    out["quotation_number"] = t["quotation_number"]
    revisions = src["quotation_number"].map(parse_revision)
    out["quotation_base_no"] = revisions.str[0]
    out["revision_no"] = revisions.str[1]
    out["odoo_rfq_number"] = t["matched_rfq_number"]
    out["odoo_order_id"] = t["matched_order_id"]
    out["revision_group"] = [
        order or (f"Q:{base}" if base else doc)
        for order, base, doc in zip(out["odoo_order_id"], out["quotation_base_no"], out["doc_id"])
    ]
    out["quotation_date"] = pd.to_datetime(t["quotation_date_final"], format="%Y-%m-%d", errors="coerce")
    out["document_date"] = pd.to_datetime(src["quotation_date_raw"].map(parse_document_date))
    out["quotation_year"] = out["quotation_date"].dt.year.astype("Int64")
    out["quotation_month"] = out["quotation_date"].dt.strftime("%Y-%m").fillna("")
    out["fiscal_year"] = out["quotation_date"].map(fiscal_year)
    out["date_source"] = t["date_source"]
    out["date_ambiguous"] = t["date_ambiguous"]

    # --- customer ---
    out["customer_id"] = t["matched_customer_id"]
    out["customer_name"] = t["matched_customer_name"]
    out["customer_type"] = t["customer_type"].str.upper()
    out["end_user_id"] = t["matched_end_user_id"]
    out["end_user_name"] = t["matched_end_user_name"]
    order_industry = t["matched_order_industry"]
    out["industry"] = order_industry.where(order_industry != "", t["matched_industry"])
    out["industry_source"] = [
        "ODOO_ORDER" if o else ("CUSTOMER_HISTORY" if c else "")
        for o, c in zip(order_industry, t["matched_industry"])
    ]
    detail = t["matched_order_industry_specify_others"].where(
        order_industry != "", t["matched_industry_specify_others"]
    )
    out["industry_detail"] = detail.str.upper().where(out["industry"].str.upper() == "OTHERS", "")
    out["city"] = t["matched_city"]
    out["state"] = src["matched_state"].map(standardize_state)
    out["country"] = t["matched_country"]
    out["sales_regions"] = src["regions"].map(standardize_regions)
    out["customer_match_quality"] = [
        match_quality(s, c) for s, c in zip(src["customer_match_status"], src["matched_customer_id"])
    ]

    # --- item ---
    out["item_no"] = t["item_no"]
    out["parent_item_no"] = t["parent_item_no"]
    out["item_level"] = pd.to_numeric(t["item_level"]).astype(int)
    out["product_name"] = src["product_name"].str.strip()
    out["description"] = src["description"].str.strip()
    out["product_family"] = t["product_family"]
    out["make_canonical"] = src["make_canonical"].str.strip()
    out["model_canonical"] = src["model_canonical"].str.strip()
    out["quantity"] = src["quantity"].map(to_number)
    out["unit_std"] = src["unit_normalized"].map(standardize_unit)
    out["unit_class"] = out["unit_std"].map(lambda u: UNIT_CLASS.get(u, ""))

    # --- pricing ---
    out["price_status"] = t["price_status"]
    out["currency_orig"] = t["currency"]
    out["currency_source"] = t["currency_source"]
    out["unit_price_orig"] = src["unit_price_final"].map(to_number)
    out["total_price_orig"] = src["total_price_final"].map(to_number)
    rates, sources = [], []
    for currency in out["currency_orig"]:
        if not currency:
            rates.append(None); sources.append("")
        elif currency == BASE_CURRENCY:
            rates.append(1.0); sources.append(FX_SOURCE_BASE)
        elif currency in FX_RATES_TO_INR:
            rates.append(FX_RATES_TO_INR[currency]); sources.append(FX_SOURCE_ODOO)
        else:
            rates.append(None); sources.append(FX_SOURCE_NO_RATE)
    out["fx_rate_to_inr"] = pd.to_numeric(pd.Series(rates, index=src.index))
    out["fx_source"] = sources
    out["unit_price_inr"] = (out["unit_price_orig"] * out["fx_rate_to_inr"]).round(2)
    out["total_price_inr"] = (out["total_price_orig"] * out["fx_rate_to_inr"]).round(2)
    out["price_basis"] = t["price_basis"]
    out["price_quality"] = t["price_quality"]
    out["item_confidence"] = t["item_confidence"]
    out["duplicate_line"] = src.duplicated(keep="first").map({True: "Y", False: "N"})

    # --- hierarchy, per document in document order ---
    line_type, parent_seq, bundle_seq = [], [], []
    for _, group in out.groupby("doc_id", sort=False):
        types, parents, bundles = classify_lines(
            group["item_level"].tolist(),
            [q is not None and not pd.isna(q) for q in group["quantity"]],
            group["price_status"].tolist(),
            group["product_name"].tolist(),
            group["item_no"].tolist(),
            group["parent_item_no"].tolist(),
        )
        line_type += types
        parent_seq += parents
        bundle_seq += bundles
    # groupby(sort=False) keeps first-appearance order and documents are
    # contiguous in the source, so the flat lists line up with out.index.
    out["line_type"] = line_type
    out["parent_line_seq"] = pd.array(parent_seq, dtype="Int64")
    out["bundle_id"] = [
        f"{doc}-{seq:04d}" if seq is not None else ""
        for doc, seq in zip(out["doc_id"], bundle_seq)
    ]

    positive = (out["total_price_orig"].fillna(0) > 0) | (out["unit_price_orig"].fillna(0) > 0)
    usable = (
        (out["price_quality"] == "TRUSTED")
        & out["price_basis"].isin(["REPORTED", "DERIVED_FROM_TOTAL"])
        & positive
        & (out["currency_orig"] != "")
        & (out["line_type"] != "TOTAL")
    )
    out["price_usable"] = usable.map({True: "Y", False: "N"})
    include = []
    for _, group in out.groupby("doc_id", sort=False):
        has_total = (group["total_price_orig"].fillna(0) > 0) & (group["price_usable"] == "Y")
        include += include_in_total(
            group["line_type"].tolist(),
            [None if pd.isna(p) else int(p) for p in group["parent_line_seq"]],
            has_total.tolist(),
        )
    out["include_in_doc_total"] = include

    # --- outcome ---
    out["quote_status"] = t["matched_order_quote_status"].str.upper().map(QUOTE_STATUS).fillna("")
    out["quote_nature"] = t["firm_or_budgetary"].str.upper()
    out["outcome_status"] = t["matched_order_present_status"].str.upper()
    out["outcome_group"] = out["outcome_status"].map(OUTCOME_GROUP).fillna("")
    out["winning_chance_pct"] = src["matched_order_winning_chance"].map(parse_percent)
    out["sbu"] = t["matched_order_sbu_type"].str.upper()
    out["spares_type"] = t["matched_order_spares_type"]
    out["service_type"] = t["matched_order_service_type"]

    # --- revisions: one latest document per revision_group ---
    docs = out.drop_duplicates("doc_id")[["doc_id", "revision_group", "revision_no", "document_date"]]
    docs = docs.sort_values(["revision_group", "revision_no", "document_date", "doc_id"], na_position="first")
    latest = set(docs.groupby("revision_group")["doc_id"].last())
    out["is_latest_revision"] = out["doc_id"].map(lambda d: "Y" if d in latest else "N")

    return out[ITEM_COLUMNS]


def build_quotations(src: pd.DataFrame, items: pd.DataFrame) -> pd.DataFrame:
    first = src.drop_duplicates("source_path").set_index("source_path")
    t = first.map(clean_text)
    doc_ids = first.index.map(doc_id_for)
    base = items.drop_duplicates("doc_id").set_index("doc_id")

    shared = [c for c in QUOTATION_COLUMNS if c in base.columns]
    q = base.loc[doc_ids, shared].copy()
    q.insert(0, "doc_id", doc_ids)

    def col(name):
        return pd.Series(t[name].values, index=doc_ids)

    q["type_of_quote"] = col("matched_order_type_of_quote").str.upper()
    q["expected_finalization"] = [
        expected_finalization(m, y) for m, y in zip(
            col("matched_order_tentative_finalization_month"), col("matched_order_finalization_year"))
    ]
    q["cycle_time_days"] = col("matched_order_average_cycle_time").map(parse_days)
    q["loss_reason"] = col("matched_order_reason_for_loss")
    q["loss_reason_detail"] = col("matched_order_main_reason_of_losing_order")
    q["po_number"] = col("matched_order_po_number")
    q["order_currency"] = col("matched_order_po_currency")
    q["order_latest_price_quoted"] = col("matched_order_latest_price_quoted")

    def money(name, zero_is_blank):
        values = col(name).map(to_number)
        return values.where(values != 0) if zero_is_blank else values

    q["order_price_inr"] = money("matched_order_price_in_inr", zero_is_blank=True)
    q["order_po_value_orig"] = money("matched_order_po_value", zero_is_blank=True)
    q["order_po_value_inr"] = money("matched_order_po_value_in_inr", zero_is_blank=True)
    q["order_potential_estimate_inr"] = money("matched_order_total_potential_estimate", zero_is_blank=True)

    included = items[items["include_in_doc_total"] == "Y"]
    q["line_count"] = items.groupby("doc_id").size().reindex(doc_ids).fillna(0).astype(int).values
    q["priced_line_count"] = included.groupby("doc_id").size().reindex(doc_ids).fillna(0).astype(int).values
    currencies = included.groupby("doc_id")["currency_orig"].agg(
        lambda s: s.iloc[0] if s.nunique() == 1 else "MIXED")
    q["doc_currency"] = currencies.reindex(doc_ids).fillna("").values
    sum_orig = included.groupby("doc_id")["total_price_orig"].sum().reindex(doc_ids)
    q["sum_items_orig"] = sum_orig.values
    q.loc[q["doc_currency"].isin(["", "MIXED"]), "sum_items_orig"] = None
    q["sum_items_inr"] = included.groupby("doc_id")["total_price_inr"].sum().reindex(doc_ids).values
    ratio = q["sum_items_inr"] / q["order_price_inr"]
    q["reconcile_ratio_vs_odoo"] = ratio.round(4)

    return q.reset_index(drop=True)[QUOTATION_COLUMNS]


def build_audit(src: pd.DataFrame, items: pd.DataFrame) -> pd.DataFrame:
    audit = src[AUDIT_SOURCE_COLUMNS].copy()
    audit.insert(0, "line_seq", items["line_seq"].values)
    audit.insert(0, "doc_id", items["doc_id"].values)
    return audit


def build_fx_table() -> pd.DataFrame:
    rows = [{"currency": BASE_CURRENCY, "rate_to_inr": 1.0, "fx_source": FX_SOURCE_BASE,
             "note": "Base currency."}]
    for currency, rate in FX_RATES_TO_INR.items():
        rows.append({"currency": currency, "rate_to_inr": rate, "fx_source": FX_SOURCE_ODOO,
                     "note": "Odoo's fixed rate, identical on every order; ignores rate movement over time."})
    return pd.DataFrame(rows)


def build_dictionary() -> pd.DataFrame:
    rows = []
    for sheet, columns in (("Items", ITEM_COLUMNS), ("Quotations", QUOTATION_COLUMNS)):
        for column in columns:
            rows.append({"sheet": sheet, "column": column, "meaning": DICTIONARY.get(column, "")})
    rows.append({"sheet": "Audit", "column": "(all)",
                 "meaning": "doc_id + line_seq + source columns moved off Items, verbatim from the knowledge bank CSV. "
                            "source_file / source_path can contain customer names."})
    return pd.DataFrame(rows)


# ============================================================
# VALIDATION
# ============================================================

def validate(src: pd.DataFrame, items: pd.DataFrame, quotations: pd.DataFrame,
             sheets: dict[str, pd.DataFrame]) -> list[tuple[str, bool, str]]:
    checks = []

    def check(name, ok, detail=""):
        checks.append((name, bool(ok), detail))

    check("row count preserved", len(items) == len(src), f"{len(src)} in, {len(items)} out")
    check("documents preserved", quotations["doc_id"].nunique() == src["source_path"].nunique() == len(quotations),
          f"{src['source_path'].nunique()} documents")
    paths = pd.Series(src["source_path"].unique())
    check("doc_id unique per source_path", paths.map(doc_id_for).nunique() == len(paths))
    orders_in = src.loc[src["matched_order_id"] != "", "matched_order_id"].nunique()
    check("orders preserved", items.loc[items["odoo_order_id"] != "", "odoo_order_id"].nunique() == orders_in,
          f"{orders_in} orders")

    source_total = src["total_price_final"].map(to_number)
    same = (source_total.fillna(-1) == items["total_price_orig"].fillna(-1)).all()
    check("total_price_orig equals source total_price_final", same)

    recomputed = (items["total_price_orig"] * items["fx_rate_to_inr"]).round(2)
    check("total_price_inr = orig x rate", (recomputed.fillna(-1) == items["total_price_inr"].fillna(-1)).all())
    check("no priced row without a rate",
          not ((items["total_price_orig"].notna()) & (items["currency_orig"] != "") & items["fx_rate_to_inr"].isna()).any(),
          f"NO_RATE rows: {(items['fx_source'] == FX_SOURCE_NO_RATE).sum()}")

    for column, allowed in ALLOWED_VALUES.items():
        bad = set(items[column].fillna("").unique()) - allowed
        check(f"allowed values: {column}", not bad, f"unexpected: {sorted(bad)}" if bad else "")

    text_columns = [c for c in ITEM_COLUMNS if is_text(items[c])]
    placeholder_hits = sum(
        items[c].astype(str).str.strip().str.upper().isin(PLACEHOLDERS - {"-"}).sum()
        for c in text_columns if c not in ("product_name", "description", "make_canonical", "model_canonical")
    )
    check("no placeholder text on standardized columns", placeholder_hits == 0, f"{placeholder_hits} cells")

    dated = items["quotation_date"].notna()
    year_ok = (items.loc[dated, "quotation_year"] == items.loc[dated, "quotation_date"].dt.year).all()
    month_ok = (items.loc[dated, "quotation_month"] == items.loc[dated, "quotation_date"].dt.strftime("%Y-%m")).all()
    check("quotation_year/month agree with quotation_date", year_ok and month_ok)
    check("quotation_date parsed wherever source had one",
          dated.sum() == (src["quotation_date_final"] != "").sum(), f"{dated.sum()} dated rows")

    latest_per_group = quotations.groupby("revision_group")["is_latest_revision"].apply(lambda s: (s == "Y").sum())
    check("exactly one latest revision per group", (latest_per_group == 1).all(),
          f"{len(latest_per_group)} groups")

    check("no order-level money on Items", not (ORDER_MONEY_COLUMNS & set(items.columns)))
    undocumented = [c for c in ITEM_COLUMNS + QUOTATION_COLUMNS if not DICTIONARY.get(c)]
    check("every column documented in Data_Dictionary", not undocumented, ", ".join(undocumented))

    linked = items[items["parent_line_seq"].notna()]
    parent_no = items.set_index(["doc_id", "line_seq"])["item_no"].map(normalize_item_no)
    keys = list(zip(linked["doc_id"], linked["parent_line_seq"].astype(int)))
    agree = (parent_no.loc[keys].values == linked["parent_item_no"].map(normalize_item_no).values).all() if keys else True
    check("every linked parent row carries the parsed parent_item_no", agree, f"{len(linked)} links")

    too_long = 0
    illegal = 0
    for frame in sheets.values():
        for c in frame.columns:
            if is_text(frame[c]):
                s = frame[c].fillna("").astype(str)
                too_long += (s.str.len() > EXCEL_CELL_LIMIT).sum()
                illegal += s.str.contains(EXCEL_ILLEGAL_RX).sum()
    check("no cell over Excel's 32,767-char limit", too_long == 0, f"{too_long} cells")
    check("no Excel-illegal control characters left", illegal == 0, f"{illegal} cells")
    return checks


def summarize(src: pd.DataFrame, items: pd.DataFrame, quotations: pd.DataFrame) -> list[str]:
    lines = []

    def counts(frame, column):
        return ", ".join(f"{k or '(blank)'}={v}" for k, v in frame[column].fillna("").value_counts().items())

    for column in ("line_type", "price_usable", "include_in_doc_total", "duplicate_line",
                   "currency_orig", "fx_source", "unit_std", "unit_class", "industry_source",
                   "customer_match_quality", "outcome_group"):
        lines.append(f"Items.{column}: {counts(items, column)}")
    lines.append(f"Quotations.is_latest_revision: {counts(quotations, 'is_latest_revision')}")
    lines.append(f"Quotations.doc_currency: {counts(quotations, 'doc_currency')}")

    subs = items[items["line_type"] == "SUB_LINE"]
    lines.append(f"SUB_LINE rows: {len(subs)}, parent printed as its own row: {subs['parent_line_seq'].notna().sum()}, "
                 f"parent not printed (e.g. 1A/1B with no row 1): {subs['parent_line_seq'].isna().sum()}")
    unparsed = items["document_date"].isna() & (src["quotation_date_raw"].str.strip() != "").values
    lines.append(f"document_date unparseable (not a date, e.g. a reference number): {unparsed.sum()} rows")

    ratio = quotations["reconcile_ratio_vs_odoo"].dropna()
    if len(ratio):
        within = ((ratio >= 0.9) & (ratio <= 1.1)).sum()
        lines.append(f"reconcile_ratio_vs_odoo: {len(ratio)} documents with both totals, "
                     f"{within} within +/-10%, median {ratio.median():.3f}")
    return lines


# ============================================================
# EXCEL
# ============================================================

def excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    for c in frame.columns:
        if is_text(frame[c]):
            frame[c] = frame[c].map(
                lambda v: EXCEL_ILLEGAL_RX.sub("", v)[:EXCEL_CELL_LIMIT] if isinstance(v, str) else v)
    return frame


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl",
                        date_format="yyyy-mm-dd", datetime_format="yyyy-mm-dd") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for position, column in enumerate(frame.columns, start=1):
                letter = get_column_letter(position)
                sample = frame[column].head(200).astype(str).str.len()
                width = max(len(column), int(sample.quantile(0.9)) if len(sample) else 0)
                sheet.column_dimensions[letter].width = min(max(width + 2, 8), 50)
                if column in MONEY_COLUMNS:
                    fmt = "#,##0.00"
                elif column in DATE_COLUMNS:
                    fmt = "yyyy-mm-dd"
                else:
                    continue
                for cell in sheet[letter][1:]:
                    cell.number_format = fmt


# ============================================================
# MAIN
# ============================================================

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--input", type=Path, help="knowledge bank CSV (default: newest dated file)")
    args = parser.parse_args()

    print("=" * 70)
    print("ANALYSIS-READY EXPORT")
    print("=" * 70)

    input_path = args.input or find_default_input()
    if input_path is None or not input_path.exists():
        print(f"\nInput not found: {input_path or INPUT_FOLDER}")
        return 1
    print(f"\nInput : {input_path}")

    src = pd.read_csv(input_path, dtype=str, keep_default_na=False)
    print(f"Rows  : {len(src)}   Documents: {src['source_path'].nunique()}")

    items = build_items(src)
    quotations = build_quotations(src, items)
    audit = build_audit(src, items)
    sheets = {
        "Items": excel_safe(items),
        "Quotations": excel_safe(quotations),
        "Audit": excel_safe(audit),
        "FX_Rates": build_fx_table(),
        "Data_Dictionary": build_dictionary(),
    }

    checks = validate(src, items, quotations, sheets)
    summary = summarize(src, items, quotations)

    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    today = date.today().isoformat()
    workbook_path = unique_output_path(OUTPUT_FOLDER, f"knowledge_bank_analysis_{today}", ".xlsx")
    log_path = workbook_path.with_name(
        workbook_path.stem.replace("knowledge_bank_analysis", "analysis_export_validation") + ".txt")

    print(f"\nWriting {workbook_path.name} ...")
    write_workbook(workbook_path, sheets)

    failed = [c for c in checks if not c[1]]
    with open(log_path, "w", encoding="utf-8") as log:
        log.write(f"Input    : {input_path}\nWorkbook : {workbook_path}\n\n")
        log.write("VALIDATION\n")
        for name, ok, detail in checks:
            log.write(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  ({detail})" if detail else "") + "\n")
        log.write("\nSUMMARY\n")
        for line in summary:
            log.write(f"  {line}\n")

    print(f"Items rows      : {len(items)}")
    print(f"Quotation rows  : {len(quotations)}")
    print(f"Checks          : {len(checks) - len(failed)} passed, {len(failed)} failed")
    print(f"Validation log  : {log_path}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
